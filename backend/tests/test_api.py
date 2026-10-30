import os
import sys
from pathlib import Path
os.environ['DATABASE_URL']='sqlite://'
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.pool import StaticPool
from sqlalchemy.orm import sessionmaker
from database import Base, get_db
from main import app, hasher, now
from models import User
from io import BytesIO
from openpyxl import Workbook

@pytest.fixture()
def client():
    engine=create_engine('sqlite://',connect_args={'check_same_thread':False},poolclass=StaticPool)
    Base.metadata.create_all(engine)
    factory=sessionmaker(bind=engine)
    def override():
        with factory() as db: yield db
    app.dependency_overrides[get_db]=override
    with factory() as db:
        for role in ['Admin','Super User','Viewer']:
            db.add(User(id=role,full_name=role,email=role.replace(' ','').lower()+'@example.com',password_hash=hasher.hash('test-password'),role=role,requested_role=role,status='Approved',registered_at=now()))
        db.commit()
    with TestClient(app) as c: yield c
    app.dependency_overrides.clear()
    engine.dispose()

def auth(c,role='Admin'):
    response=c.post('/api/auth/signin',json={'email':role.replace(' ','').lower()+'@example.com','password':'test-password','requestedRole':role})
    assert response.status_code==200
    return {'Authorization':'Bearer '+response.json()['token']}

def record(c,h):
    response=c.post('/api/fines',headers=h,json={'customerId':'C001','fineAmount':'125.50','fineDate':'2026-10-01','reasonCategory':'Missing proof','sourceArea':'Warehouse'})
    assert response.status_code==201,response.text
    assert response.json()['currencyCode']=='USD'
    return response.json()

def test_roles_signup_approval_logout(client):
    c=client
    assert c.get('/api/fines').status_code==401
    body={'fullName':'New User','email':'new@example.com','password':'password123','requestedRole':'Super User'}
    assert c.post('/api/auth/signup',json=body).status_code==201
    assert c.post('/api/auth/signup',json=body).status_code==409
    login={'email':body['email'],'password':body['password'],'requestedRole':'Super User'}
    assert c.post('/api/auth/signin',json=login).json()['detail']=='pending'
    viewer=auth(c,'Viewer')
    assert c.get('/api/users/requests',headers=viewer).status_code==403
    admin=auth(c)
    requests=c.get('/api/users/requests',headers=admin).json()
    user=next(u for u in requests if u['email']==body['email'])
    assert c.post('/api/users/'+user['id']+'/decision',headers=admin,json={'decision':'approve','assignedRole':'Super User'}).status_code==200
    assert c.post('/api/auth/signin',json=login).status_code==200
    assert c.post('/api/auth/logout',headers=viewer).status_code==200
    assert c.get('/api/fines',headers=viewer).status_code==401

def test_reviews_filters_kpis_and_concurrency(client):
    c=client; h=auth(c); f=record(c,h)
    viewer=auth(c,'Viewer'); reviewer=auth(c,'Super User')
    body={'decision':'Invalid','comment':'Proof does not support deduction','expectedStatus':'Pending','expectedVersion':1}
    path='/api/fines/'+f['fineId']+'/reviews'
    assert c.post(path,headers=viewer,json=body).status_code==403
    assert c.post(path,headers=reviewer,json={**body,'comment':'  '}).status_code==422
    updated=c.post(path,headers=reviewer,json=body)
    assert updated.status_code==200,updated.text
    assert updated.json()['version']==2
    assert c.post(path,headers=reviewer,json=body).status_code==409
    assert len(c.get(path,headers=h).json())==1
    assert c.get('/api/fines?status=Pending',headers=h).json()==[]
    assert len(c.get('/api/fines?sourceArea=Warehouse&search=C001',headers=h).json())==1
    assert c.get('/api/fines?fromDate=2026-11-01&toDate=2026-01-01',headers=h).status_code==422
    assert float(c.get('/api/dashboard',headers=h).json()['deductionsBySource']['Warehouse'])==125.5
    exported=c.post('/api/exports/fines',headers=viewer,json={'ids':[f['fineId']]})
    assert exported.status_code==200
    from openpyxl import load_workbook
    book=load_workbook(BytesIO(exported.content))
    assert book.active.max_row==2
    assert book.active.cell(2,6).value==125.5
    assert book.active.cell(2,10).value=='Invalid'
    assert book.active.cell(2,7).value=='USD'

def workbook(amount):
    w=Workbook(); w.active.append(['CustomerID','Fines Imposed','Status','Reason','Date','SourceArea']);w.active.append(['C002',amount,'Pending','Missing evidence','2026-10-02','Logistics']);b=BytesIO();w.save(b);return b.getvalue()

def test_atomic_import_and_duplicate_protection(client):
    c=client; h=auth(c)
    def upload(data,commit=False): return c.post('/api/imports/validate?commit='+str(commit).lower(),headers=h,files={'file':('data.xlsx',data,'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')})
    assert upload(workbook(-1),True).json()['valid'] is False
    assert c.get('/api/fines',headers=h).json()==[]
    data=workbook(220)
    assert upload(data).json()['imported']==0
    assert upload(data,True).json()['imported']==1
    assert upload(data,True).status_code==409
    assert len(c.get('/api/fines',headers=h).json())==1
    assert c.post('/api/imports/validate',headers=h,files={'file':('bad.xlsx',b'broken')}).status_code==422


def test_usd_contract(client):
    c=client; h=auth(c)
    assert c.get('/api/dashboard',headers=h).json()['currencyCode']=='USD'
    body={'customerId':'C001','fineAmount':'125.50','fineDate':'2026-10-01','reasonCategory':'Missing proof','sourceArea':'Warehouse','currencyCode':'INR'}
    assert c.post('/api/fines',headers=h,json=body).status_code==422

@pytest.mark.parametrize('role,summary,manage,review,import_status', [
    ('Viewer',403,403,403,403),
    ('Super User',200,403,200,403),
    ('Admin',200,200,200,200),
])
def test_role_page_and_api_matrix(client,role,summary,manage,review,import_status):
    c=client; admin=auth(c); fine=record(c,admin); headers=auth(c,role)
    assert c.get('/api/fines',headers=headers).status_code==200
    assert c.get('/api/dashboard',headers=headers).status_code==summary
    assert c.get('/api/users/requests',headers=headers).status_code==manage
    assert c.post('/api/users/nonexistent/decision',headers=headers,json={'decision':'approve'}).status_code==(409 if role=='Admin' else 403)
    assert c.post('/api/fines/'+fine['fineId']+'/reviews',headers=headers,json={'decision':'Valid','comment':'Evidence verified','expectedStatus':'Pending','expectedVersion':1}).status_code==review
    assert c.post('/api/imports/validate',headers=headers,files={'file':('data.xlsx',workbook(100))}).status_code==import_status
    assert c.post('/api/fines',headers=headers,json={'customerId':'C002','fineAmount':'125.50','fineDate':'2026-10-01','reasonCategory':'Missing proof','sourceArea':'Warehouse'}).status_code==(201 if role=='Admin' else 403)

def test_registration_cannot_request_admin(client):
    assert client.post('/api/auth/signup',json={'fullName':'Escalation','email':'escalate@example.com','password':'password123','requestedRole':'Admin'}).status_code==422

def test_rejection_wrong_role_and_expired_sessions(client):
    c=client; admin=auth(c)
    body={'fullName':'Rejected User','email':'rejected@example.com','password':'password123','requestedRole':'Viewer'}
    c.post('/api/auth/signup',json=body).raise_for_status()
    user=next(u for u in c.get('/api/users/requests',headers=admin).json() if u['email']==body['email'])
    path='/api/users/'+user['id']+'/decision'
    assert c.post(path,headers=admin,json={'decision':'reject'}).status_code==200
    assert c.post(path,headers=admin,json={'decision':'approve'}).status_code==409
    assert c.post('/api/auth/signin',json={'email':body['email'],'password':body['password'],'requestedRole':'Viewer'}).json()['detail']=='rejected'
    assert c.post('/api/auth/signin',json={'email':'admin@example.com','password':'test-password','requestedRole':'Viewer'}).json()['detail']=='wrong-role'
    assert c.get('/api/fines',headers={'Authorization':'Bearer invalid'}).status_code==401
    from models import Session
    from sqlalchemy import update
    generator=app.dependency_overrides[get_db]()
    db=next(generator)
    db.execute(update(Session).values(expires_at='2000-01-01T00:00:00+00:00')); db.commit(); generator.close()
    assert c.get('/api/fines',headers=admin).status_code==401

def test_filters_and_export_preserve_literal_text(client):
    c=client; h=auth(c)
    payload={'customerId':'0002','customerName':'=HYPERLINK("example")','fineAmount':'0.10','fineDate':'2026-10-10','reasonCategory':'Packaging','sourceArea':'Manufacturing'}
    first=c.post('/api/fines',headers=h,json=payload).json()
    second=c.post('/api/fines',headers=h,json={**payload,'fineAmount':'0.20','fineDate':'2026-10-11','sourceArea':'Logistics'}).json()
    result=c.get('/api/fines?customerId=0002&sourceArea=Manufacturing&reason=Packaging&fromDate=2026-10-10&toDate=2026-10-10',headers=h)
    assert [f['fineId'] for f in result.json()]==[first['fineId']]
    summary=c.get('/api/dashboard',headers=h).json()
    assert str(summary['totalDeductions'])=='0.3' or str(summary['totalDeductions'])=='0.30'
    from openpyxl import load_workbook
    exported=c.post('/api/exports/fines',headers=h,json={'ids':[first['fineId'],first['fineId'],second['fineId']]}).content
    book=load_workbook(BytesIO(exported))
    assert book.active.max_row==3
    assert book.active.cell(2,3).data_type=='s'
    assert book.active.cell(2,3).value.startswith('=HYPERLINK')

def test_workbook_validation_boundaries(client):
    c=client; h=auth(c)
    def send(content,name='data.xlsx'):
        return c.post('/api/imports/validate?commit=true',headers=h,files={'file':(name,content)})
    assert send(b'x','data.csv').status_code==422
    assert send(b'x'*(5*1024*1024+1)).status_code==413
    book=Workbook(); book.active.append(['Wrong headers']); data=BytesIO(); book.save(data)
    assert send(data.getvalue()).status_code==422
    book=Workbook(); book.active.append(['CustomerID','Fines Imposed','Status','Reason','Date','SourceArea'])
    book.active.append(['0001',100,'Pending','Evidence','2026-10-02','Warehouse'])
    book.active.append(['0002',-1,'Pending','Evidence','bad-date','Warehouse'])
    data=BytesIO(); book.save(data)
    result=send(data.getvalue()).json()
    assert result['valid'] is False and result['errors'][0]['row']==3
    assert c.get('/api/fines',headers=h).json()==[]

def test_sqlite_persistence_after_reconnect(client,tmp_path):
    from models import Fine, Review
    from main import make_fine, FineInput
    database_path='sqlite:///'+str(tmp_path/'persist.db')
    engine=create_engine(database_path)
    Base.metadata.create_all(engine)
    with sessionmaker(bind=engine)() as db:
        fine=make_fine(FineInput(customerId='0005',fineAmount='42.75',fineDate='2026-10-15',reasonCategory='Evidence',sourceArea='Warehouse'))
        db.add(fine); db.commit(); fine_id=fine.id
    engine.dispose()
    reopened=create_engine(database_path)
    with sessionmaker(bind=reopened)() as db:
        fine=db.get(Fine,fine_id)
        assert str(fine.amount)=='42.75' and fine.data['currencyCode']=='USD'
    reopened.dispose()
