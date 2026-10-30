import hashlib
import io
import os
import secrets
from contextlib import asynccontextmanager
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal
from typing import Literal, Optional
from uuid import uuid4

from fastapi import FastAPI, Depends, HTTPException, UploadFile, File, Query
from fastapi.middleware.cors import CORSMiddleware
from fastapi.security import HTTPBearer, HTTPAuthorizationCredentials
from openpyxl import load_workbook
from pydantic import BaseModel, EmailStr, Field, ValidationError, field_validator
from pwdlib import PasswordHash
from sqlalchemy import select, update, text
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session as DBSession
from database import Base, engine, get_db
from models import User, Session, Fine, Review

Status = Literal['Valid', 'Invalid', 'Pending']
Role = Literal['Admin', 'Super User', 'Viewer']
Source = Literal['Logistics', 'Warehouse', 'Manufacturing']
hasher = PasswordHash.recommended()
bearer = HTTPBearer(auto_error=False)
def now(): return datetime.now(timezone.utc).isoformat()
def uid(): return str(uuid4())
def fail(code, detail): raise HTTPException(code, detail)
def public(user): return {'id': user.id, 'fullName': user.full_name, 'email': user.email, 'role': user.role}

@asynccontextmanager
async def lifespan(app):
    Base.metadata.create_all(engine)
    yield

app = FastAPI(title='Fines & Fees API', version='2.0.0', lifespan=lifespan)
app.add_middleware(CORSMiddleware, allow_origins=os.getenv('FRONTEND_ORIGIN', 'http://localhost:5173').split(','), allow_methods=['GET','POST','PATCH'], allow_headers=['Authorization','Content-Type'])

def current(credentials: Optional[HTTPAuthorizationCredentials] = Depends(bearer), db: DBSession = Depends(get_db)):
    if not credentials: fail(401, 'unauthorized')
    session = db.get(Session, hashlib.sha256(credentials.credentials.encode()).hexdigest())
    if not session or session.expires_at < now(): fail(401, 'session-expired')
    user = db.get(User, session.user_id)
    if not user or user.status != 'Approved' or user.role not in ['Admin', 'Super User', 'Viewer']: fail(403, 'unauthorized')
    return user

def admin(user=Depends(current)):
    if user.role != 'Admin': fail(403, 'unauthorized')
    return user

def reviewer(user=Depends(current)):
    if user.role not in ['Admin','Super User']: fail(403, 'unauthorized')
    return user

class Registration(BaseModel):
    fullName: str = Field(min_length=1, max_length=120)
    email: EmailStr
    password: str = Field(min_length=8, max_length=128)
    requestedRole: Literal['Super User','Viewer'] = 'Viewer'
    @field_validator('fullName')
    @classmethod
    def name(cls, value):
        if not value.strip(): raise ValueError('Name is required')
        return value.strip()

class Login(BaseModel):
    email: EmailStr
    password: str = Field(max_length=128)
    requestedRole: Role

@app.get('/api/health')
def health(): return {'status':'ok'}

@app.get('/api/health/database')
def database_health(user=Depends(admin), db: DBSession=Depends(get_db)):
    db.execute(text('SELECT 1'))
    return {'status':'connected', 'database':engine.dialect.name}

@app.post('/api/auth/signup', status_code=201)
def signup(body: Registration, db: DBSession=Depends(get_db)):
    user = User(id=uid(), full_name=body.fullName, email=str(body.email).lower(), password_hash=hasher.hash(body.password), role=body.requestedRole, requested_role=body.requestedRole, status='Pending Approval', registered_at=now())
    db.add(user)
    try: db.commit()
    except IntegrityError:
        db.rollback(); fail(409, 'account-exists')
    return {'status':user.status}

@app.post('/api/auth/signin')
def signin(body: Login, db: DBSession=Depends(get_db)):
    user = db.scalar(select(User).where(User.email == str(body.email).lower()))
    if not user or not hasher.verify(body.password, user.password_hash): fail(401, 'invalid-credentials')
    if user.status != 'Approved': fail(403, 'pending' if user.status == 'Pending Approval' else 'rejected')
    if user.role != body.requestedRole: fail(403, 'wrong-role')
    token = secrets.token_urlsafe(48)
    db.add(Session(token_hash=hashlib.sha256(token.encode()).hexdigest(), user_id=user.id, expires_at=(datetime.now(timezone.utc)+timedelta(hours=8)).isoformat()))
    db.commit()
    return {'token':token, 'user':public(user)}

@app.get('/api/auth/me')
def me(user=Depends(current)): return public(user)

@app.post('/api/auth/logout')
def logout(credentials: HTTPAuthorizationCredentials=Depends(bearer), user=Depends(current), db: DBSession=Depends(get_db)):
    db.delete(db.get(Session, hashlib.sha256(credentials.credentials.encode()).hexdigest())); db.commit()
    return {'status':'signed-out'}

@app.get('/api/users/requests')
def requests(user=Depends(admin), db: DBSession=Depends(get_db)):
    return [dict(public(u), requestedRole=u.requested_role, assignedRole=u.role if u.status=='Approved' else None, status=u.status, registeredAt=u.registered_at) for u in db.scalars(select(User).order_by(User.registered_at.desc()))]

class Approval(BaseModel):
    decision: Literal['approve','reject']
    assignedRole: Literal['Super User','Viewer'] = 'Viewer'

@app.post('/api/users/{user_id}/decision')
def decide(user_id: str, body: Approval, user=Depends(admin), db: DBSession=Depends(get_db)):
    result = db.execute(update(User).where(User.id==user_id, User.status=='Pending Approval').values(status='Approved' if body.decision=='approve' else 'Rejected', role=body.assignedRole))
    if result.rowcount != 1: fail(409, 'conflict')
    db.commit(); return {'status':'saved'}

class FineInput(BaseModel):
    customerId: str = Field(min_length=1, max_length=120)
    customerName: str = Field(default='', max_length=120)
    fineAmount: Decimal = Field(ge=0, max_digits=18, decimal_places=2)
    status: Status = 'Pending'
    reasonCategory: str = Field(min_length=1, max_length=200)
    fineDate: date
    sourceArea: Source
    orderReference: str = Field(default='', max_length=120)
    currencyCode: Literal['USD'] = 'USD'
    @field_validator('customerId', 'reasonCategory')
    @classmethod
    def nonblank(cls, value):
        if not value.strip(): raise ValueError('Must not be blank')
        return value.strip()

def output(fine): return dict(fine.data, fineId=fine.id, status=fine.status, fineAmount=float(fine.amount), version=fine.version)
def make_fine(body):
    data=body.model_dump(mode='json')
    data.update(createdAt=now(), updatedAt=now())
    return Fine(id=uid(), customer_id=body.customerId, amount=body.fineAmount, status=body.status, fine_date=body.fineDate.isoformat(), source=body.sourceArea, version=1, data=data)

def filtered(db, status=None, customerId=None, sourceArea=None, fromDate=None, toDate=None, reason=None, search=None):
    if fromDate and toDate and fromDate > toDate: fail(422, 'invalid-date-range')
    query=select(Fine)
    for column, value in [(Fine.status,status),(Fine.customer_id,customerId),(Fine.source,sourceArea)]:
        if value: query=query.where(column==value)
    if fromDate: query=query.where(Fine.fine_date>=fromDate.isoformat())
    if toDate: query=query.where(Fine.fine_date<=toDate.isoformat())
    records=[output(f) for f in db.scalars(query.order_by(Fine.fine_date.desc(), Fine.id))]
    if reason: records=[f for f in records if f['reasonCategory']==reason]
    if search:
        records=[f for f in records if search.lower() in ' '.join(str(f.get(k,'')) for k in ['fineId','customerId','customerName','orderReference']).lower()]
    return records

@app.get('/api/fines')
def fines(status: Optional[Status]=None, customerId: Optional[str]=None, sourceArea: Optional[Source]=None, fromDate: Optional[date]=None, toDate: Optional[date]=None, reason: Optional[str]=None, search: Optional[str]=Query(None,max_length=200), user=Depends(current), db: DBSession=Depends(get_db)):
    return filtered(db,status,customerId,sourceArea,fromDate,toDate,reason,search)

@app.get('/api/dashboard')
def dashboard(user=Depends(reviewer), db: DBSession=Depends(get_db)):
    rows=filtered(db)
    return {'totalFines':len(rows), 'totalDeductions':sum(Decimal(str(f['fineAmount'])) for f in rows), 'pendingCases':sum(f['status']=='Pending' for f in rows), 'deductionsBySource':{s:sum(Decimal(str(f['fineAmount'])) for f in rows if f['sourceArea']==s) for s in ['Logistics','Warehouse','Manufacturing']}, 'currencyCode':'USD'}

@app.post('/api/fines',status_code=201)
def create(body: FineInput, user=Depends(admin), db: DBSession=Depends(get_db)):
    fine=make_fine(body); db.add(fine); db.commit(); return output(fine)

@app.get('/api/fines/{fine_id}/reviews')
def reviews(fine_id: str, user=Depends(current), db: DBSession=Depends(get_db)):
    if not db.get(Fine,fine_id): fail(404,'not-found')
    return sorted([r.data for r in db.scalars(select(Review).where(Review.fine_id==fine_id))],key=lambda r:r['reviewedAt'],reverse=True)

class Decision(BaseModel):
    decision: Status
    comment: str = Field(min_length=1,max_length=1000)
    expectedStatus: Status
    expectedVersion: int = Field(ge=1)
    @field_validator('comment')
    @classmethod
    def comment_required(cls,value):
        if not value.strip(): raise ValueError('Comment is required')
        return value.strip()

@app.post('/api/fines/{fine_id}/reviews')
def review(fine_id: str, body: Decision, user=Depends(reviewer), db: DBSession=Depends(get_db)):
    fine=db.get(Fine,fine_id)
    if not fine: fail(404,'not-found')
    stamp=now()
    data=dict(fine.data,updatedAt=stamp,lastReviewedAt=stamp,lastReviewerName=user.full_name,lastReviewComment=body.comment)
    result=db.execute(update(Fine).where(Fine.id==fine_id,Fine.version==body.expectedVersion,Fine.status==body.expectedStatus).values(status=body.decision,version=Fine.version+1,data=data))
    if result.rowcount!=1: db.rollback(); fail(409,'conflict')
    review_id=uid()
    db.add(Review(id=review_id,fine_id=fine_id,data={'reviewId':review_id,'fineId':fine_id,'oldStatus':body.expectedStatus,'newStatus':body.decision,'reviewerId':user.id,'reviewerName':user.full_name,'comment':body.comment,'reviewedAt':stamp}))
    db.commit(); db.refresh(fine); return output(fine)

@app.post('/api/imports/validate')
async def import_excel(file: UploadFile=File(...), commit: bool=False, user=Depends(admin), db: DBSession=Depends(get_db)):
    """Validate the complete workbook before any inserts. Explicit commit is required."""
    if not file.filename or not file.filename.lower().endswith('.xlsx'): fail(422,'xlsx-required')
    content=await file.read(5*1024*1024+1)
    if len(content)>5*1024*1024: fail(413,'file-too-large')
    try:
        book=load_workbook(io.BytesIO(content),read_only=True,data_only=True)
        sheet=book.active
        if sheet.max_row and sheet.max_row>10001: fail(422,'maximum-10000-rows')
        rows=sheet.iter_rows(values_only=True)
        headers=next(rows)
        required=['CustomerID','Fines Imposed','Status','Reason','Date','SourceArea']
        if len(set(headers)) != len(headers) or any(h not in headers for h in required): fail(422,{'requiredColumns':required})
        valid=[]; errors=[]
        for number,row in enumerate(rows,2):
            if number>10001: fail(422,'maximum-10000-rows')
            if all(v is None for v in row): continue
            values=dict(zip(headers,row))
            try:
                raw_date=values['Date']
                if isinstance(raw_date,datetime): raw_date=raw_date.date()
                valid.append(FineInput(customerId=str(values['CustomerID']) if values['CustomerID'] is not None else '',fineAmount=values['Fines Imposed'],status=values['Status'],reasonCategory=values['Reason'],fineDate=raw_date,sourceArea=values['SourceArea']))
            except (ValidationError,ValueError,TypeError) as error:
                errors.append({'row':number,'message':str(error)})
        book.close()
    except HTTPException: raise
    except Exception: fail(422,'invalid-workbook')
    if errors: return {'valid':False,'imported':0,'errors':errors}
    if not valid: fail(422,'empty-workbook')
    if commit:
        # A content digest prevents the same file from being imported twice.
        digest=hashlib.sha256(content).hexdigest()
        db.add(ImportBatch(id=digest))
        db.add_all([make_fine(row) for row in valid])
        try: db.commit()
        except IntegrityError: db.rollback(); fail(409,'already-imported')
    return {'valid':True,'rows':len(valid),'imported':len(valid) if commit else 0,'errors':[]}

from sqlalchemy import Column, String
class ImportBatch(Base):
    __tablename__='import_batches'
    id=Column(String(64),primary_key=True)

class ExportRequest(BaseModel):
    ids: list[str] = Field(min_length=1,max_length=10000)

@app.post('/api/exports/fines')
def export_fines(body: ExportRequest, user=Depends(current), db: DBSession=Depends(get_db)):
    from openpyxl import Workbook
    from fastapi.responses import StreamingResponse
    book=Workbook(); sheet=book.active; sheet.title='Fine Records'
    sheet.append(['Fine ID','Customer ID','Customer','Order Reference','Date','Fines Imposed','Currency','Reason','Source Area','Status','Latest Comment'])
    # Chunk IDs to stay within SQL Server's parameter limit.
    ids=list(dict.fromkeys(body.ids))
    for offset in range(0,len(ids),500):
        for fine in db.scalars(select(Fine).where(Fine.id.in_(ids[offset:offset+500]))):
            f=output(fine)
            sheet.append([f['fineId'],f['customerId'],f.get('customerName',''),f.get('orderReference',''),f['fineDate'],float(fine.amount),'USD',f['reasonCategory'],f['sourceArea'],f['status'],f.get('lastReviewComment','')])
            for cell in sheet[sheet.max_row]:
                if isinstance(cell.value,str): cell.data_type='s'
            sheet.cell(sheet.max_row,6).number_format='#,##0.00'
    sheet.freeze_panes='A2'; sheet.auto_filter.ref=sheet.dimensions
    from openpyxl.utils import get_column_letter
    for index in range(1,12): sheet.column_dimensions[get_column_letter(index)].width=24
    buffer=io.BytesIO(); book.save(buffer); buffer.seek(0)
    return StreamingResponse(buffer,media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':'attachment; filename="fines_export.xlsx"'})
